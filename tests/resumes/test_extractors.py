from __future__ import annotations

from io import BytesIO

import pytest


def test_txt_extraction_normalizes_bom_nul_and_line_endings():
    """Catches parser input corruption from common exported-text artifacts."""
    from app.resumes.extractors import extract_text

    text = extract_text("resume.txt", b"\xef\xbb\xbfPython\x00\r\nFastAPI")
    assert text == "Python\nFastAPI"


def test_docx_extraction_reads_paragraphs():
    """Catches accepting DOCX while silently producing no profile text."""
    from docx import Document

    from app.resumes.extractors import extract_text

    document = Document()
    document.add_paragraph("Python 后端工程师")
    document.add_paragraph("使用 FastAPI 构建服务")
    buffer = BytesIO()
    document.save(buffer)

    assert extract_text("resume.docx", buffer.getvalue()) == "Python 后端工程师\n使用 FastAPI 构建服务"


def test_extractor_rejects_unsupported_or_oversized_file():
    """Catches unbounded or executable content reaching document parsers."""
    from app.core.exceptions import UnsupportedFileError
    from app.resumes.extractors import FilePolicy, extract_text

    with pytest.raises(UnsupportedFileError):
        extract_text("resume.html", b"<script>bad</script>")
    with pytest.raises(UnsupportedFileError, match="10 MiB"):
        FilePolicy().validate("large.txt", "text/plain", b"x" * (10 * 1024 * 1024 + 1))


@pytest.mark.parametrize("filename", ["broken.pdf", "broken.docx"])
def test_malformed_supported_document_returns_stable_ingestion_error(filename):
    """Catches third-party parser exceptions escaping the API contract."""
    from app.core.exceptions import IngestionError
    from app.resumes.extractors import extract_text

    with pytest.raises(IngestionError) as captured:
        extract_text(filename, b"not a real document")
    assert captured.value.code == "resume_extraction_failed"


@pytest.mark.parametrize(
    "filename,media_type,content",
    [
        ("fake.pdf", "application/pdf", b"plain text"),
        ("fake.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", b"plain text"),
        ("fake.txt", "text/plain", b"%PDF-1.7\nnot text"),
    ],
)
def test_upload_rejects_misleading_signature(filename, media_type, content):
    from app.core.exceptions import IngestionError
    from app.resumes.extractors import FilePolicy

    with pytest.raises(IngestionError):
        FilePolicy().validate(filename, media_type, content)


def test_docx_expansion_bomb_rejected_before_document_parser():
    from zipfile import ZIP_DEFLATED, ZipFile
    from app.core.exceptions import IngestionError
    from app.resumes.extractors import FilePolicy

    stream = BytesIO()
    with ZipFile(stream, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"a" * (32 * 1024 * 1024))
    assert len(stream.getvalue()) < 10 * 1024 * 1024
    with pytest.raises(IngestionError) as captured:
        FilePolicy().validate("bomb.docx", "application/octet-stream", stream.getvalue())
    assert captured.value.code == "document_resource_limit"


@pytest.mark.parametrize("kind", ["pages", "content", "decompression"])
def test_pdf_resource_limits_apply_before_extraction(kind):
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, NameObject
    from app.core.exceptions import IngestionError
    from app.resumes.extractors import FilePolicy

    writer = PdfWriter()
    page = writer.add_blank_page(100, 100)
    if kind == "pages":
        for _ in range(100):
            writer.add_blank_page(100, 100)
    else:
        stream = DecodedStreamObject()
        stream.set_data(b" " * ((3 if kind == "content" else 9) * 1024 * 1024))
        page[NameObject("/Contents")] = writer._add_object(stream.flate_encode())
    output = BytesIO()
    writer.write(output)
    assert len(output.getvalue()) < 10 * 1024 * 1024
    with pytest.raises(IngestionError) as error:
        FilePolicy().validate("budget.pdf", "application/pdf", output.getvalue())
    assert error.value.code in {"document_resource_limit", "resume_extraction_failed"}


def test_preparation_reads_chunks_and_closes_temporary_stream_on_exception():
    from app.resumes.extractors import FilePolicy

    class ChunkSource(BytesIO):
        def read(self, size=-1):
            assert 0 < size <= 65536
            return super().read(size)

    held = None
    with pytest.raises(RuntimeError, match="synthetic"):
        with FilePolicy().prepare("resume.txt", "text/plain", ChunkSource(b"Python")) as prepared:
            held = prepared.stream
            assert prepared.sha256 == "18885f27b5af9012df19e496460f9294d5ab76128824c6f993787004f6d9a7db"
            raise RuntimeError("synthetic")
    assert held.closed
